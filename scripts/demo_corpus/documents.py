"""确定性 PDF / DOCX / XLSX 生成组件。

字节可复现的关键点：
- PDF 使用 reportlab invariant 模式固定时间戳，并显式固定文档 ID；
- DOCX / XLSX 生成后统一重写 OOXML ZIP：固定条目顺序、时间戳、权限与压缩参数；
- openpyxl 会在保存时写入当前时间，因此重写 ZIP 时强制替换 ``dcterms:modified``。
"""

from __future__ import annotations

import hashlib
import re
import zipfile
from pathlib import Path

from docx import Document
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from reportlab import rl_config

rl_config.invariant = 1

from reportlab.lib import colors  # noqa: E402
from reportlab.lib.pagesizes import A4  # noqa: E402
from reportlab.pdfbase import pdfmetrics  # noqa: E402
from reportlab.pdfbase.ttfonts import TTFont  # noqa: E402
from reportlab.pdfgen import canvas  # noqa: E402

from . import facts  # noqa: E402

PAGE_WIDTH, PAGE_HEIGHT = A4
MARGIN_LEFT = 56.0
MARGIN_RIGHT = 56.0
MARGIN_TOP = 64.0
MARGIN_BOTTOM = 60.0
CONTENT_WIDTH = PAGE_WIDTH - MARGIN_LEFT - MARGIN_RIGHT

FIXED_ZIP_DATE = (1980, 1, 1, 0, 0, 0)
ZIP_COMPRESS_LEVEL = 9

PDF_FONT_NAME = "DemoSC"
PDF_FONT_NOTE = "NotoSansSC-VF.ttf"


# ---------------------------------------------------------------------------
# OOXML ZIP 归一化
# ---------------------------------------------------------------------------


def normalize_ooxml_zip(source: Path, target: Path, core_modified: str | None = None) -> None:
    """按固定顺序、时间戳、权限和压缩参数重写 OOXML（zip）容器。"""
    with zipfile.ZipFile(source) as archive:
        names = sorted(archive.namelist())
        payload = {name: archive.read(name) for name in names}

    if core_modified is not None and "docProps/core.xml" in payload:
        text = payload["docProps/core.xml"].decode("utf-8")
        text = re.sub(
            r"(<dcterms:modified[^>]*>)[^<]*(</dcterms:modified>)",
            lambda match: match.group(1) + core_modified + match.group(2),
            text,
        )
        payload["docProps/core.xml"] = text.encode("utf-8")

    with zipfile.ZipFile(target, "w", zipfile.ZIP_DEFLATED, compresslevel=ZIP_COMPRESS_LEVEL) as archive:
        for name in names:
            info = zipfile.ZipInfo(name, date_time=FIXED_ZIP_DATE)
            info.compress_type = zipfile.ZIP_DEFLATED
            info.create_system = 3
            info.external_attr = 0o100644 << 16
            info.internal_attr = 0
            info.extra = b""
            info.comment = b""
            archive.writestr(info, payload[name])


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------


def register_pdf_font(font_path: Path) -> str:
    if PDF_FONT_NAME not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont(PDF_FONT_NAME, str(font_path)))
    return PDF_FONT_NAME


class PdfWriter:
    """手工排版的 PDF 写入器：页码与章节页码在渲染时确定，便于生成定位信息。"""

    def __init__(self, path: Path, *, title: str, subject: str, keywords: str, document_key: str):
        self.path = Path(path)
        self.canvas = canvas.Canvas(str(self.path), pagesize=A4, invariant=1)
        self.canvas.setTitle(title)
        self.canvas.setAuthor(facts.DOC_AUTHOR)
        self.canvas.setSubject(subject)
        self.canvas.setCreator(facts.DOC_PRODUCER)
        self.canvas.setKeywords(keywords)
        self.page = 1
        self.y = PAGE_HEIGHT - MARGIN_TOP
        self.sections: list[str] = []
        self.section_pages: dict[str, int] = {}
        self._draw_footer()
        self._document_key = document_key

    # -- 基础绘制 ---------------------------------------------------------

    def _draw_footer(self) -> None:
        self.canvas.setFont(PDF_FONT_NAME, 8.5)
        self.canvas.setFillColor(colors.HexColor("#666666"))
        self.canvas.drawString(
            MARGIN_LEFT,
            MARGIN_BOTTOM - 22,
            f"{facts.FICTION_MARKER} · 第 {self.page} 页",
        )
        self.canvas.setFillColor(colors.black)

    def _new_page(self) -> None:
        self.canvas.showPage()
        self.page += 1
        self.y = PAGE_HEIGHT - MARGIN_TOP
        self._draw_footer()

    def _ensure(self, height: float) -> None:
        if self.y - height < MARGIN_BOTTOM:
            self._new_page()

    def _wrap(self, text: str, size: float) -> list[str]:
        lines: list[str] = []
        current = ""
        for char in text:
            if char == "\n":
                lines.append(current)
                current = ""
                continue
            trial = current + char
            if current and pdfmetrics.stringWidth(trial, PDF_FONT_NAME, size) > CONTENT_WIDTH:
                lines.append(current)
                current = char
            else:
                current = trial
        lines.append(current)
        return lines

    def _paragraph(self, text: str, size: float, leading: float, indent: float = 0.0) -> None:
        for line in self._wrap(text, size):
            self._ensure(leading)
            self.canvas.setFont(PDF_FONT_NAME, size)
            self.canvas.drawString(MARGIN_LEFT + indent, self.y - size, line)
            self.y -= leading

    # -- 内容块 -----------------------------------------------------------

    def banner(self, text: str) -> None:
        size = 11.0
        leading = 16.0
        lines = self._wrap(text, size)
        height = leading * len(lines) + 10
        self._ensure(height + 8)
        top = self.y
        self.canvas.setStrokeColor(colors.HexColor("#888888"))
        self.canvas.rect(MARGIN_LEFT, top - height, CONTENT_WIDTH, height, stroke=1, fill=0)
        self.canvas.setFont(PDF_FONT_NAME, size)
        cursor = top - 5
        for line in lines:
            self.canvas.drawString(MARGIN_LEFT + 6, cursor - size, line)
            cursor -= leading
        self.y = top - height - 12

    def title(self, text: str) -> None:
        size = 16.0
        leading = 24.0
        self._ensure(leading + 8)
        self._paragraph(text, size, leading)
        self.y -= 6

    def h1(self, text: str) -> None:
        size = 12.5
        leading = 20.0
        self._ensure(leading + 10)
        self.y -= 6
        self._paragraph(text, size, leading)
        self.sections.append(text)
        self.section_pages[text] = self.page

    def h2(self, text: str) -> None:
        size = 11.0
        leading = 17.0
        self._ensure(leading + 6)
        self.y -= 3
        self._paragraph(text, size, leading)

    def body(self, text: str) -> None:
        self._paragraph(text, 10.5, 16.0)

    def bullet(self, text: str) -> None:
        self._paragraph("· " + text, 10.5, 16.0, indent=6.0)

    def note(self, text: str) -> None:
        self._paragraph(text, 9.0, 14.0)

    def spacer(self, height: float = 8.0) -> None:
        self.y -= height

    def page_break(self) -> None:
        self._new_page()

    def table(self, headers: list[str], rows: list[list[str]], widths: list[float]) -> None:
        size = 9.0
        leading = 13.5
        scale = CONTENT_WIDTH / sum(widths)
        columns = [w * scale for w in widths]

        def draw_row(cells: list[str], bold: bool) -> None:
            wrapped = [
                self._wrap_cell(str(cell), size, columns[index] - 8) for index, cell in enumerate(cells)
            ]
            row_height = leading * max(len(item) for item in wrapped)
            self._ensure(row_height + 2)
            top = self.y
            x = MARGIN_LEFT
            self.canvas.setFont(PDF_FONT_NAME, size)
            for index, lines in enumerate(wrapped):
                for offset, line in enumerate(lines):
                    self.canvas.drawString(x + 4, top - size - 2 - offset * leading, line)
                x += columns[index]
            self.y = top - row_height - 4
            self.canvas.setStrokeColor(colors.HexColor("#cccccc"))
            self.canvas.line(MARGIN_LEFT, self.y + 2, MARGIN_LEFT + CONTENT_WIDTH, self.y + 2)
            if bold:
                self.canvas.setStrokeColor(colors.HexColor("#888888"))
                self.canvas.line(MARGIN_LEFT, self.y + 2, MARGIN_LEFT + CONTENT_WIDTH, self.y + 2)

        draw_row(headers, True)
        for row in rows:
            draw_row(row, False)
        self.y -= 6

    def _wrap_cell(self, text: str, size: float, width: float) -> list[str]:
        lines: list[str] = []
        current = ""
        for char in text:
            trial = current + char
            if current and pdfmetrics.stringWidth(trial, PDF_FONT_NAME, size) > width:
                lines.append(current)
                current = char
            else:
                current = trial
        lines.append(current)
        return lines

    # -- 收尾 -------------------------------------------------------------

    def finish(self) -> dict:
        self.canvas._doc._ID = (
            b"\n[<" + self._document_key.encode("ascii") + b"><" + self._document_key.encode("ascii") + b">]\n"
            b"% ReportLab generated PDF document -- digest (http://www.reportlab.com)\n"
        )
        self.canvas.showPage()
        self.canvas.save()
        return {"page_count": self.page, "sections": list(self.sections), "section_pages": dict(self.section_pages)}


def pdf_document_key(file_name: str) -> str:
    seed = f"{facts.DATASET_NAME}|{facts.DATASET_VERSION}|{facts.GENERATOR_VERSION}|{file_name}"
    return hashlib.md5(seed.encode("utf-8")).hexdigest()


# ---------------------------------------------------------------------------
# DOCX
# ---------------------------------------------------------------------------


def build_docx(path: Path, *, title: str, subject: str, blocks: list[tuple]) -> dict:
    tmp = Path(str(path) + ".tmp")
    document = Document()
    core = document.core_properties
    core.author = facts.DOC_AUTHOR
    core.last_modified_by = facts.DOC_AUTHOR
    core.created = _fixed_datetime()
    core.modified = _fixed_datetime()
    core.title = title
    core.subject = subject
    core.category = "demo-corpus"
    core.comments = facts.FICTION_NOTICE
    core.revision = 1

    sections: list[str] = []
    for block in blocks:
        kind = block[0]
        if kind == "marker":
            paragraph = document.add_paragraph()
            run = paragraph.add_run(block[1])
            run.bold = True
        elif kind == "h1":
            document.add_heading(block[1], level=1)
            sections.append(block[1])
        elif kind == "h2":
            document.add_heading(block[1], level=2)
        elif kind == "body":
            document.add_paragraph(block[1])
        elif kind == "note":
            paragraph = document.add_paragraph()
            paragraph.add_run(block[1]).italic = True
        elif kind == "bullet":
            document.add_paragraph(block[1], style="List Bullet")
        elif kind == "table":
            _add_docx_table(document, block[1], block[2])
        else:  # pragma: no cover - 防御性分支
            raise ValueError(f"unknown docx block: {kind}")

    document.save(tmp)
    normalize_ooxml_zip(tmp, path)
    tmp.unlink()
    return {"sections": sections}


def _add_docx_table(document, headers: list[str], rows: list[list[str]]) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for index, header in enumerate(headers):
        table.rows[0].cells[index].text = header
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cells[index].text = str(value)


# ---------------------------------------------------------------------------
# XLSX
# ---------------------------------------------------------------------------


def build_xlsx(path: Path, *, title: str, subject: str, sheets: list[tuple[str, list[list]]]) -> dict:
    tmp = Path(str(path) + ".tmp")
    workbook = Workbook()
    workbook.remove(workbook.active)
    sheet_rows: dict[str, int] = {}
    for name, rows in sheets:
        worksheet = workbook.create_sheet(title=name)
        for row in rows:
            worksheet.append(row)
        worksheet["A1"].font = Font(bold=True)
        _autosize(worksheet, rows)
        sheet_rows[name] = worksheet.max_row

    props = workbook.properties
    props.creator = facts.DOC_AUTHOR
    props.lastModifiedBy = facts.DOC_AUTHOR
    props.created = _fixed_datetime()
    props.modified = _fixed_datetime()
    props.title = title
    props.subject = subject
    props.category = "demo-corpus"
    props.description = facts.FICTION_NOTICE

    workbook.save(tmp)
    normalize_ooxml_zip(tmp, path, core_modified=facts.GENERATED_AT)
    tmp.unlink()
    names = list(sheet_rows)
    return {"sheet_names": names, "sheet_rows": sheet_rows}


def _autosize(worksheet, rows: list[list]) -> None:
    if not rows:
        return
    width_count = max(len(row) for row in rows)
    for index in range(width_count):
        width = 8
        for row in rows:
            if index < len(row) and row[index] is not None:
                width = max(width, _display_width(str(row[index])) + 3)
        worksheet.column_dimensions[get_column_letter(index + 1)].width = min(width, 42)


def _display_width(text: str) -> int:
    return sum(2 if ord(char) > 0x2E80 else 1 for char in text)


def _fixed_datetime():
    from datetime import datetime

    return datetime(2026, 9, 25, 0, 0, 0)
