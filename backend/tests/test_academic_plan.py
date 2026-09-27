"""阶段 7C：``POST /api/academic/plan`` 契约、可见性、证据映射与确定性测试。

全部离线：不访问网络、不调用 LLM / Embedding / Reranker、不下载模型。
"""

from __future__ import annotations

import json
import sys
from pathlib import Path

import openpyxl
import pytest

from app.academic.types import (
    WARN_CATEGORY_MISMATCH,
    WARN_COURSE_NOT_IN_RULE,
    WARN_RECORD_CONTRADICTION,
    WARN_TIME_CONFLICT,
    WARN_VERSION_CONFLICT,
)
from app.models import AcademicRecordSet, AcademicRuleSet, Document, DocumentChunk
from tests.conftest import demo_file

PLAN_ENDPOINT = "/api/academic/plan"
RECORDS_ENDPOINT = "/api/academic/records/import"
RULES_ENDPOINT = "/api/academic/rules/import"
OPTIONS_ENDPOINT = "/api/academic/options"
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

SPEC_FIELDS = {
    "required_credits",
    "completed_credits",
    "in_progress_credits",
    "remaining_credits",
    "missing_required_courses",
    "category_gaps",
    "conflict_warnings",
    "evidence",
}
EVIDENCE_FIELDS = {
    "chunk_id",
    "doc_id",
    "file_name",
    "document_version",
    "effective_from",
    "page_number",
    "sheet_name",
    "row_start",
    "row_end",
    "section_title",
    "quote",
}


# ---------------------------------------------------------------------------
# 夹具与辅助
# ---------------------------------------------------------------------------


@pytest.fixture
def demo_ready(ingest_demo, context, worker):
    """已投影 active demo 学业资料（Student A/B + 2025.1/2026.1）。"""
    ingest_demo()
    return context


@pytest.fixture
def db(context):
    """不含任何 demo 数据的孤立数据库（用于 upload 通道的确定性断言）。"""
    return context


def _assert_error(response, code: str) -> dict:
    assert response.status_code >= 400, response.text
    payload = response.json()
    assert payload["code"] == code, payload
    assert payload["request_id"], "所有后端生成的错误都必须带非空 request_id"
    assert set(payload) == {"code", "message", "details", "request_id"}
    assert "Traceback" not in response.text
    assert "/app/" not in response.text
    return payload


def _options(client) -> dict:
    response = client.get(OPTIONS_ENDPOINT)
    assert response.status_code == 200, response.text
    return response.json()


def _record_set_id(client, name: str) -> str:
    for item in _options(client)["record_sets"]:
        if item["name"] == name:
            return item["id"]
    raise AssertionError(f"未找到课程记录集合 {name}")


def _rule_set_id(client, version: str) -> str:
    for item in _options(client)["rule_sets"]:
        if item["rule_version"] == version:
            return item["id"]
    raise AssertionError(f"未找到培养方案 {version}")


def _plan(client, record_set_id: str, rule_set_id: str):
    return client.post(
        PLAN_ENDPOINT, json={"record_set_id": record_set_id, "rule_set_id": rule_set_id}
    )


def _plan_ok(client, record_set_id: str, rule_set_id: str) -> dict:
    response = _plan(client, record_set_id, rule_set_id)
    assert response.status_code == 200, response.text
    return response.json()


def _import(client, endpoint: str, path: Path, *, mime: str | None = None) -> str:
    with open(path, "rb") as handle:
        files = {"file": (path.name, handle, mime or "application/octet-stream")}
        response = client.post(endpoint, files=files)
    assert response.status_code == 200, response.text
    return response.json()["id"]


def _write_records_xlsx(tmp_path: Path, rows, *, header=None) -> Path:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "课程记录"
    sheet.append(
        list(header or ("序号", "课程代码", "课程名称", "学分", "课程类别", "状态", "学期"))
    )
    for row in rows:
        sheet.append(list(row))
    path = tmp_path / "records.xlsx"
    workbook.save(str(path))
    return path


def _write_plan_xlsx(tmp_path: Path, *, total="160.0 学分", minimums=None, catalog=None) -> Path:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "培养方案"
    sheet.append(["专业名称", "招生年份", "文档版本", "生效日期", "毕业总学分"])
    sheet.append(["计算机科学与技术", 2026, "2026.9", "2026-09-01", total])

    minimum_sheet = workbook.create_sheet("类别最低学分")
    minimum_sheet.append(["课程类别", "最低学分"])
    for category, credits in minimums if minimums is not None else (("公共必修", 52), ("专业必修", 58)):
        minimum_sheet.append([category, credits])

    catalog_sheet = workbook.create_sheet("课程目录")
    catalog_sheet.append(["课程代码", "课程名称", "学分", "课程类别"])
    for row in catalog or (
        ("QM-CS101", "程序设计基础", 4, "专业必修"),
        ("QM-CS102", "高等数学（一）", 5, "公共必修"),
    ):
        catalog_sheet.append(list(row))

    path = tmp_path / "plan.xlsx"
    workbook.save(str(path))
    return path


def _warnings(payload: dict) -> dict[str, dict]:
    return {item["code"]: item for item in payload["conflict_warnings"]}


def _assert_evidence_integrity(context, payload: dict) -> None:
    evidence = payload["evidence"]
    pairs = [(item["doc_id"], item["chunk_id"]) for item in evidence]
    assert pairs == sorted(pairs), "evidence 必须按 (doc_id, chunk_id) 稳定排序"
    assert all(item["chunk_id"] for item in evidence), "不允许空字符串 chunk ID"
    assert len(pairs) == len(set(pairs)), "evidence 必须去重"

    with context.session_factory() as session:
        for item in evidence:
            chunk = session.get(DocumentChunk, item["chunk_id"])
            assert chunk is not None, "evidence.chunk_id 必须对应真实 DocumentChunk 行"
            assert chunk.doc_id == item["doc_id"], "doc_id 必须与 chunk 所属文档一致"
            document = session.get(Document, chunk.doc_id)
            assert document is not None
            assert item["file_name"] == document.file_name
            assert item["quote"] == chunk.text, "quote 必须是被引用 chunk 的真实文本"
            locator = chunk.locator
            assert item["sheet_name"] == locator.get("sheet_name")
            assert item["row_start"] == locator.get("row_start")
            assert item["row_end"] == locator.get("row_end")
            assert item["page_number"] == locator.get("page_number")
            assert item["section_title"] == locator.get("section_title")

    known = {item["chunk_id"] for item in evidence}
    for warning in payload["conflict_warnings"]:
        assert set(warning["evidence_chunk_ids"]) <= known, warning["code"]
    for item in payload["missing_required_courses"]:
        assert set(item["evidence_chunk_ids"]) <= known, item["course_code"]


# ---------------------------------------------------------------------------
# 1—3：请求契约与可见性
# ---------------------------------------------------------------------------


def test_plan_request_body_is_strictly_two_ids(client, demo_ready) -> None:
    record_set_id = _record_set_id(client, "匿名学生A · 课程记录")
    rule_set_id = _rule_set_id(client, "2026.1")
    assert _plan(client, record_set_id, rule_set_id).status_code == 200

    extra = client.post(
        PLAN_ENDPOINT,
        json={
            "record_set_id": record_set_id,
            "rule_set_id": rule_set_id,
            "major": "计算机科学与技术",
        },
    )
    assert extra.status_code == 422
    assert extra.json()["request_id"]
    assert extra.json()["code"] == "REQUEST_VALIDATION_ERROR"

    assert client.post(PLAN_ENDPOINT, json={"record_set_id": record_set_id}).status_code == 422
    assert client.post(PLAN_ENDPOINT, json={"rule_set_id": rule_set_id}).status_code == 422
    assert client.post(PLAN_ENDPOINT, json={}).status_code == 422
    assert (
        client.post(
            PLAN_ENDPOINT, json={"record_set_id": "not-a-uuid", "rule_set_id": rule_set_id}
        ).status_code
        == 422
    )
    assert (
        client.post(
            PLAN_ENDPOINT, json={"record_set_id": record_set_id, "rule_set_id": "1234"}
        ).status_code
        == 422
    )


def test_unknown_ids_return_stable_errors(client, demo_ready) -> None:
    record_set_id = _record_set_id(client, "匿名学生A · 课程记录")
    rule_set_id = _rule_set_id(client, "2026.1")
    missing_id = "00000000-0000-4000-8000-000000000000"

    _assert_error(_plan(client, missing_id, rule_set_id), "ACADEMIC_RECORD_SET_NOT_FOUND")
    _assert_error(_plan(client, record_set_id, missing_id), "ACADEMIC_RULE_SET_NOT_FOUND")


def test_invisible_demo_sets_are_rejected(client, demo_ready) -> None:
    """inactive / candidate / 旧版本 demo 集合都不可选。"""
    payloads: dict[str, str] = {}
    with demo_ready.session_factory() as session:
        for state, version in (
            ("inactive", "2026.1"),
            ("candidate", "2026.1"),
            ("active", "2019.0"),
        ):
            document = Document(
                source_type="demo",
                source_key=f"demo:probe:{state}:{version}",
                file_name=f"probe-{state}-{version}.pdf",
                file_type="pdf",
                mime_type="application/pdf",
                sha256=(state + version).ljust(64, "0")[:64],
                doc_category="degree_plan",
                dataset_version=version,
                status="ready",
                retrievable=True,
                activation_state=state,
            )
            session.add(document)
            session.flush()
            rule_set = AcademicRuleSet(
                source_type="demo",
                source_key=document.source_key,
                dataset_version=version,
                activation_state=state,
                status="ready",
                display_name=f"probe-{state}-{version}",
                major="计算机科学与技术",
                admission_year=2025,
                rule_version=f"probe-{state}",
                required_credits="150.0",
                content_hash=(state + version).ljust(64, "a")[:64],
                source_doc_id=document.id,
            )
            session.add(rule_set)
            session.flush()
            payloads[f"{state}:{version}"] = rule_set.id
        session.commit()

    record_set_id = _record_set_id(client, "匿名学生A · 课程记录")
    assert len(payloads) == 3
    for rule_set_id in payloads.values():
        _assert_error(_plan(client, record_set_id, rule_set_id), "ACADEMIC_RULE_SET_NOT_FOUND")


def test_deleted_source_document_is_rejected(client, db) -> None:
    record_set_id = _import(client, RECORDS_ENDPOINT, demo_file("13-课程记录-匿名学生A"), mime=XLSX_MIME)
    rule_set_id = _import(client, RULES_ENDPOINT, demo_file("01-培养方案"), mime="application/pdf")

    with db.session_factory() as session:
        document_id = session.get(AcademicRecordSet, record_set_id).source_doc_id
    assert client.delete(f"/api/documents/{document_id}").status_code == 204

    _assert_error(_plan(client, record_set_id, rule_set_id), "ACADEMIC_SOURCE_INVALID")


# ---------------------------------------------------------------------------
# 4—10：显式选择、确定性数字与类别规则
# ---------------------------------------------------------------------------


def test_rule_version_is_never_selected_implicitly(client, demo_ready) -> None:
    record_set_id = _record_set_id(client, "匿名学生A · 课程记录")
    newest = _plan_ok(client, record_set_id, _rule_set_id(client, "2026.1"))
    older = _plan_ok(client, record_set_id, _rule_set_id(client, "2025.1"))

    assert newest["required_credits"] == 160.0
    assert older["required_credits"] == 155.0
    assert set(newest) == SPEC_FIELDS
    for forbidden in ("rule_version", "major", "record_set_id", "rule_set_id"):
        assert forbidden not in newest


@pytest.mark.parametrize(
    "version, required, extra_missing",
    [("2025.1", 155.0, None), ("2026.1", 160.0, "QM-CS303")],
)
def test_demo_student_a_plan(
    client, demo_ready, version: str, required: float, extra_missing: str | None
) -> None:
    payload = _plan_ok(
        client, _record_set_id(client, "匿名学生A · 课程记录"), _rule_set_id(client, version)
    )

    assert payload["required_credits"] == required
    assert payload["completed_credits"] == 20.5
    assert payload["in_progress_credits"] == 9.0
    assert payload["remaining_credits"] == round(required - 20.5 - 9.0, 1)
    codes = [item["course_code"] for item in payload["missing_required_courses"]]
    expected = ["QM-CS204", "QM-CS301", "QM-CS302"]
    if extra_missing is not None:
        expected.append(extra_missing)
    assert codes == expected
    _assert_evidence_integrity(demo_ready, payload)


def test_demo_student_b_plan(client, demo_ready) -> None:
    payload = _plan_ok(
        client, _record_set_id(client, "匿名学生B · 课程记录"), _rule_set_id(client, "2026.1")
    )

    assert payload["required_credits"] == 160.0
    assert payload["completed_credits"] == 14.0
    assert payload["in_progress_credits"] == 10.5
    assert payload["remaining_credits"] == 135.5
    assert [item["course_code"] for item in payload["missing_required_courses"]] == [
        "QM-CS105",
        "QM-CS204",
        "QM-CS301",
        "QM-CS302",
        "QM-CS303",
    ]
    # Student B 的 QM-CS105 两条记录都是「不及格」，不构成矛盾
    assert WARN_RECORD_CONTRADICTION not in _warnings(payload)
    _assert_evidence_integrity(demo_ready, payload)


def test_upload_records_and_rules_plan(client, demo_ready) -> None:
    record_set_id = _import(client, RECORDS_ENDPOINT, demo_file("13-课程记录-匿名学生A"), mime=XLSX_MIME)
    rule_set_id = _import(client, RULES_ENDPOINT, demo_file("01-培养方案"), mime="application/pdf")

    payload = _plan_ok(client, record_set_id, rule_set_id)
    assert payload["required_credits"] == 155.0
    assert payload["completed_credits"] == 20.5
    assert payload["in_progress_credits"] == 9.0
    assert payload["remaining_credits"] == 125.5
    _assert_evidence_integrity(demo_ready, payload)


def test_mixed_demo_and_upload_selection(client, demo_ready) -> None:
    upload_records = _import(
        client, RECORDS_ENDPOINT, demo_file("13-课程记录-匿名学生A"), mime=XLSX_MIME
    )
    upload_rules = _import(client, RULES_ENDPOINT, demo_file("02-培养方案"), mime="application/pdf")

    cross = _plan_ok(client, upload_records, _rule_set_id(client, "2025.1"))
    assert cross["required_credits"] == 155.0
    reverse = _plan_ok(client, _record_set_id(client, "匿名学生A · 课程记录"), upload_rules)
    assert reverse["required_credits"] == 160.0
    assert reverse["remaining_credits"] == 130.5
    _assert_evidence_integrity(demo_ready, reverse)


def test_retake_duplicate_failed_and_in_progress(client, db, tmp_path) -> None:
    """重修只计一次；failed 不计分；在修单独统计；passed+in_progress 产生矛盾告警。"""
    records = _write_records_xlsx(
        tmp_path,
        rows=[
            (1, "QM-CS101", "程序设计基础", "4.0", "专业必修", "不及格", "2025-2026-1"),
            (2, "QM-CS101", "程序设计基础", "4.0", "专业必修", "通过", "2025-2026-2"),
            (3, "QM-CS101", "程序设计基础", "4.0", "专业必修", "在修", "2026-2027-1"),
            (4, "QM-CS102", "高等数学（一）", "5.0", "公共必修", "不及格", "2025-2026-1"),
        ],
    )
    plan_path = _write_plan_xlsx(tmp_path, total="10.0 学分", minimums=(("公共必修", 5), ("专业必修", 4)))
    record_set_id = _import(client, RECORDS_ENDPOINT, records, mime=XLSX_MIME)
    rule_set_id = _import(client, RULES_ENDPOINT, plan_path, mime=XLSX_MIME)

    payload = _plan_ok(client, record_set_id, rule_set_id)
    # QM-CS101 只计一次 4.0（passed 优先），failed 的 QM-CS102 不计分
    assert payload["completed_credits"] == 4.0
    assert payload["in_progress_credits"] == 0.0
    assert payload["remaining_credits"] == 6.0
    assert WARN_RECORD_CONTRADICTION in _warnings(payload)
    _assert_evidence_integrity(db, payload)


def test_category_priority_and_warnings(client, db, tmp_path) -> None:
    """类别以所选规则为准；记录类别不一致或不在目录时产生稳定 warning。"""
    records = _write_records_xlsx(
        tmp_path,
        rows=[
            (1, "QM-CS101", "程序设计基础", "4.0", "公共必修", "通过", "2025-2026-1"),
            (2, "QM-XX999", "目录外课程", "3.0", "专业选修", "通过", "2025-2026-2"),
        ],
    )
    plan_path = _write_plan_xlsx(
        tmp_path, total="10.0 学分", minimums=(("公共必修", 5), ("专业必修", 4))
    )
    record_set_id = _import(client, RECORDS_ENDPOINT, records, mime=XLSX_MIME)
    rule_set_id = _import(client, RULES_ENDPOINT, plan_path, mime=XLSX_MIME)

    payload = _plan_ok(client, record_set_id, rule_set_id)
    warnings = _warnings(payload)
    assert WARN_CATEGORY_MISMATCH in warnings
    assert WARN_COURSE_NOT_IN_RULE in warnings
    gaps = {item["category"]: item for item in payload["category_gaps"]}
    assert gaps["专业必修"]["completed_credits"] == 4.0
    _assert_evidence_integrity(db, payload)


def test_zero_remaining_still_reports_missing_required(client, db, tmp_path) -> None:
    records = _write_records_xlsx(
        tmp_path, rows=[(1, "QM-XX999", "目录外课程", "3.0", "专业必修", "通过", "2025-2026-1")]
    )
    plan_path = _write_plan_xlsx(
        tmp_path,
        total="0.0 学分",
        minimums=(("专业必修", 0),),
        catalog=(("QM-CS101", "程序设计基础", 4, "专业必修"),),
    )
    record_set_id = _import(client, RECORDS_ENDPOINT, records, mime=XLSX_MIME)
    rule_set_id = _import(client, RULES_ENDPOINT, plan_path, mime=XLSX_MIME)

    payload = _plan_ok(client, record_set_id, rule_set_id)
    assert payload["remaining_credits"] == 0.0
    assert [item["course_code"] for item in payload["missing_required_courses"]] == ["QM-CS101"]


# ---------------------------------------------------------------------------
# 11—15：冲突与证据
# ---------------------------------------------------------------------------


def test_version_conflict_covers_both_versions(client, demo_ready) -> None:
    payload = _plan_ok(
        client, _record_set_id(client, "匿名学生A · 课程记录"), _rule_set_id(client, "2026.1")
    )
    warning = _warnings(payload)[WARN_VERSION_CONFLICT]
    assert warning["severity"] == "warning"
    assert "2025.1" in warning["message"] and "2026.1" in warning["message"]

    evidence = {item["chunk_id"]: item for item in payload["evidence"]}
    names = {evidence[chunk_id]["file_name"] for chunk_id in warning["evidence_chunk_ids"]}
    assert any("2026" in name for name in names), "必须覆盖所选版本来源"
    assert any("2025" in name for name in names), "必须覆盖冲突版本来源"
    _assert_evidence_integrity(demo_ready, payload)


def test_record_contradiction_covers_record_evidence(client, demo_ready) -> None:
    payload = _plan_ok(
        client, _record_set_id(client, "匿名学生A · 课程记录"), _rule_set_id(client, "2026.1")
    )
    warning = _warnings(payload)[WARN_RECORD_CONTRADICTION]
    assert "QM-CS102" in warning["message"]
    assert warning["evidence_chunk_ids"], "矛盾告警必须带真实证据"
    evidence = {item["chunk_id"]: item for item in payload["evidence"]}
    record_doc = evidence[warning["evidence_chunk_ids"][0]]["doc_id"]
    assert all(
        evidence[chunk_id]["doc_id"] == record_doc for chunk_id in warning["evidence_chunk_ids"]
    )
    _assert_evidence_integrity(demo_ready, payload)


def test_time_conflict_covers_both_courses(client, demo_ready) -> None:
    payload = _plan_ok(
        client, _record_set_id(client, "匿名学生A · 课程记录"), _rule_set_id(client, "2026.1")
    )
    warning = _warnings(payload)[WARN_TIME_CONFLICT]
    evidence = {item["chunk_id"]: item for item in payload["evidence"]}
    quotes = "".join(evidence[chunk_id]["quote"] for chunk_id in warning["evidence_chunk_ids"])
    assert "QM-GE101" in quotes and "QM-CS201" in quotes, "证据必须覆盖冲突双方课程来源"
    _assert_evidence_integrity(demo_ready, payload)


def test_record_schedule_conflict_requires_provable_overlap(client, db, tmp_path) -> None:
    """只有同单位、同星期且区间真实重叠才产生时间冲突；无法解析即不猜。"""
    plan_path = _write_plan_xlsx(
        tmp_path, total="10.0 学分", minimums=(("公共必修", 5), ("专业必修", 4))
    )
    rule_set_id = _import(client, RULES_ENDPOINT, plan_path, mime=XLSX_MIME)
    header = ("序号", "课程代码", "课程名称", "学分", "课程类别", "状态", "学期", "上课时间")

    different_units = _write_records_xlsx(
        tmp_path,
        header=header,
        rows=[
            (1, "QM-CS101", "程序设计基础", "4.0", "专业必修", "在修", "2026-2027-1", "星期二第3-4节"),
            (2, "QM-CS102", "高等数学（一）", "5.0", "公共必修", "在修", "2026-2027-1", "星期二 10:00-11:40"),
        ],
    )
    payload = _plan_ok(client, _import(client, RECORDS_ENDPOINT, different_units, mime=XLSX_MIME), rule_set_id)
    assert WARN_TIME_CONFLICT not in _warnings(payload), "不同单位不得猜测重叠"

    overlapping = _write_records_xlsx(
        tmp_path,
        header=header,
        rows=[
            (1, "QM-CS101", "程序设计基础", "4.0", "专业必修", "在修", "2026-2027-1", "星期二 10:00-11:40"),
            (2, "QM-CS102", "高等数学（一）", "5.0", "公共必修", "在修", "2026-2027-1", "星期二 09:30-10:30"),
        ],
    )
    payload2 = _plan_ok(
        client, _import(client, RECORDS_ENDPOINT, overlapping, mime=XLSX_MIME), rule_set_id
    )
    warning = _warnings(payload2)[WARN_TIME_CONFLICT]
    evidence = {item["chunk_id"]: item for item in payload2["evidence"]}
    quotes = "".join(evidence[chunk_id]["quote"] for chunk_id in warning["evidence_chunk_ids"])
    assert "QM-CS101" in quotes and "QM-CS102" in quotes
    _assert_evidence_integrity(db, payload2)


# ---------------------------------------------------------------------------
# 17—19 / 26—29：证据真实性、不泄漏、确定性、隔离
# ---------------------------------------------------------------------------


def test_plan_response_is_strict_and_does_not_leak_internals(client, demo_ready) -> None:
    response = _plan(
        client, _record_set_id(client, "匿名学生A · 课程记录"), _rule_set_id(client, "2026.1")
    )
    body = response.text
    for leaked in (
        "source_key",
        "storage_path",
        "safe_storage_name",
        "sha256",
        "/app/",
        "chunker_fingerprint",
        "embedding_fingerprint",
        "dense_score",
        "fused_score",
        "rerank_score",
    ):
        assert leaked not in body, leaked

    payload = response.json()
    assert set(payload) == SPEC_FIELDS
    for item in payload["evidence"]:
        assert set(item) == EVIDENCE_FIELDS
    for item in payload["category_gaps"]:
        assert set(item) == {
            "category",
            "required_credits",
            "completed_credits",
            "in_progress_credits",
            "remaining_credits",
        }
    for item in payload["missing_required_courses"]:
        assert set(item) == {
            "course_code",
            "course_name",
            "credits",
            "category",
            "evidence_chunk_ids",
        }
    for item in payload["conflict_warnings"]:
        assert set(item) == {"code", "message", "severity", "evidence_chunk_ids"}


def test_plan_numbers_are_non_negative_and_one_decimal(client, demo_ready) -> None:
    payload = _plan_ok(
        client, _record_set_id(client, "匿名学生B · 课程记录"), _rule_set_id(client, "2025.1")
    )
    for key in (
        "required_credits",
        "completed_credits",
        "in_progress_credits",
        "remaining_credits",
    ):
        assert payload[key] >= 0
        assert round(payload[key], 1) == payload[key]
    for gap in payload["category_gaps"]:
        assert gap["remaining_credits"] >= 0
        for key in (
            "required_credits",
            "completed_credits",
            "in_progress_credits",
            "remaining_credits",
        ):
            assert round(gap[key], 1) == gap[key]


def test_plan_is_byte_stable_and_restart_safe(client, context, demo_ready) -> None:
    from app.academic.planning import build_plan
    from app.academic.types import planning_result_payload
    from app.db import create_db_engine, create_session_factory

    record_set_id = _record_set_id(client, "匿名学生A · 课程记录")
    rule_set_id = _rule_set_id(client, "2026.1")
    first = _plan(client, record_set_id, rule_set_id)
    second = _plan(client, record_set_id, rule_set_id)
    assert first.text == second.text, "相同请求必须字节级稳定"

    engine = create_db_engine(context.settings)
    with create_session_factory(engine)() as session:
        payload = planning_result_payload(
            build_plan(session, context.settings, record_set_id, rule_set_id)
        )
    engine.dispose()
    assert json.dumps(payload, ensure_ascii=False, sort_keys=True) == json.dumps(
        first.json(), ensure_ascii=False, sort_keys=True
    )


def test_planning_path_never_touches_models(client, demo_ready, monkeypatch) -> None:
    context = client.app.state.context

    def _forbidden(*_args, **_kwargs):  # pragma: no cover - 命中即失败
        raise AssertionError("规划路径不得访问模型或网络")

    monkeypatch.setattr(type(context.embeddings), "embed_documents", _forbidden, raising=False)
    monkeypatch.setattr(type(context.reranker), "rerank", _forbidden, raising=False)
    monkeypatch.setattr(type(context.llm), "chat", _forbidden, raising=False)

    payload = _plan_ok(
        client, _record_set_id(client, "匿名学生A · 课程记录"), _rule_set_id(client, "2026.1")
    )
    assert payload["required_credits"] == 160.0


def test_planning_modules_do_not_import_model_or_network_layers() -> None:
    package_root = Path(sys.modules["app"].__file__).resolve().parent
    forbidden = ("app.llm", "app.embedding", "app.rerank", "httpx", "requests", "socket", "urllib")
    for relative in ("academic/planning.py", "academic/engine.py", "academic/types.py"):
        source = (package_root / relative).read_text(encoding="utf-8")
        for token in forbidden:
            assert token not in source, f"{relative} 不得依赖 {token}"


def test_health_reports_planning_ready(client) -> None:
    payload = client.get("/api/health").json()
    assert payload["capabilities"]["planning"] == "ready"
    assert payload["status"] == "degraded"
    assert isinstance(payload["version"], str) and payload["version"]
