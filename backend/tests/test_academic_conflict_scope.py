"""阶段 7C 独立回归修复轮：时间冲突的范围与区间边界（BUG-7C-01）。

不变量：

- 时间冲突只针对**所选记录集中 ``status=in_progress`` 的课程**；
- 必须按 ``semester`` 隔离，缺失或无法匹配学期时不猜测；
- active demo 课表只能为「所选记录中同学期的在修课程」提供冲突证据，
  不得把与所选记录无关的课表冲突附加到结果；
- 时钟区间按**半开区间**判断（首尾相接不算重叠），节次是**端点包含**的离散区间
  （共享节次算冲突）。
"""

from __future__ import annotations

from pathlib import Path

import openpyxl
import pytest

from app.academic.types import WARN_TIME_CONFLICT

PLAN_ENDPOINT = "/api/academic/plan"
RECORDS_ENDPOINT = "/api/academic/records/import"
RULES_ENDPOINT = "/api/academic/rules/import"
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

RECORD_HEADER = (
    "序号",
    "课程代码",
    "课程名称",
    "学分",
    "课程类别",
    "状态",
    "学期",
    "上课时间",
)


@pytest.fixture
def demo_ready(ingest_demo, context, worker):
    ingest_demo()
    return context


@pytest.fixture
def db(context):
    """不含任何 demo 数据的孤立数据库（用于 upload 通道的确定性断言）。"""
    return context


def _write_records_xlsx(tmp_path: Path, rows, *, header=RECORD_HEADER, name="records.xlsx") -> Path:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "课程记录"
    sheet.append(list(header))
    for row in rows:
        sheet.append(list(row))
    path = tmp_path / name
    workbook.save(str(path))
    return path


def _write_plan_xlsx(tmp_path: Path, *, total="160.0 学分") -> Path:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "培养方案"
    sheet.append(["专业名称", "招生年份", "文档版本", "生效日期", "毕业总学分"])
    sheet.append(["计算机科学与技术", 2026, "2026.9", "2026-09-01", total])

    minimum_sheet = workbook.create_sheet("类别最低学分")
    minimum_sheet.append(["课程类别", "最低学分"])
    minimum_sheet.append(["公共必修", 5])
    minimum_sheet.append(["专业必修", 4])

    catalog_sheet = workbook.create_sheet("课程目录")
    catalog_sheet.append(["课程代码", "课程名称", "学分", "课程类别"])
    catalog_sheet.append(["QM-CS101", "程序设计基础", 4, "专业必修"])
    catalog_sheet.append(["QM-CS102", "高等数学（一）", 5, "公共必修"])

    path = tmp_path / "plan.xlsx"
    workbook.save(str(path))
    return path


def _import(client, endpoint: str, path: Path, *, mime: str | None = None) -> str:
    with open(path, "rb") as handle:
        files = {"file": (path.name, handle, mime or "application/octet-stream")}
        response = client.post(endpoint, files=files)
    assert response.status_code == 200, response.text
    return response.json()["id"]


def _plan(client, record_set_id: str, rule_set_id: str) -> dict:
    response = client.post(
        PLAN_ENDPOINT, json={"record_set_id": record_set_id, "rule_set_id": rule_set_id}
    )
    assert response.status_code == 200, response.text
    return response.json()


def _warnings(payload: dict) -> dict[str, dict]:
    return {item["code"]: item for item in payload["conflict_warnings"]}


def _demo_ids(client) -> tuple[dict[str, str], dict[str, str]]:
    options = client.get("/api/academic/options").json()
    return (
        {item["name"]: item["id"] for item in options["record_sets"]},
        {item["rule_version"]: item["id"] for item in options["rule_sets"]},
    )


def _upload_rule_set(client, tmp_path: Path) -> str:
    return _import(client, RULES_ENDPOINT, _write_plan_xlsx(tmp_path), mime=XLSX_MIME)


# ---------------------------------------------------------------------------
# 范围：只针对所选记录中同学期的在修课程
# ---------------------------------------------------------------------------


def test_schedule_conflict_is_ignored_when_unrelated_to_selected_records(
    client, demo_ready, tmp_path
) -> None:
    """demo 课表里确实存在冲突，但所选记录与该冲突课程无关 -> 不得产生时间冲突。"""
    records = _write_records_xlsx(
        tmp_path,
        rows=[
            (1, "QM-XX001", "无关课程一", "4.0", "专业必修", "在修", "2026-2027-1", None),
            (2, "QM-XX002", "无关课程二", "3.0", "公共必修", "在修", "2026-2027-1", None),
        ],
    )
    payload = _plan(
        client,
        _import(client, RECORDS_ENDPOINT, records, mime=XLSX_MIME),
        _upload_rule_set(client, tmp_path),
    )
    assert WARN_TIME_CONFLICT not in _warnings(payload)


def test_schedule_conflict_is_reported_for_related_in_progress_course(
    client, demo_ready, tmp_path
) -> None:
    """所选记录包含课表冲突的一方在修课程 -> 冲突成立，且证据覆盖双方课程来源。"""
    records = _write_records_xlsx(
        tmp_path,
        rows=[(1, "QM-CS201", "数据结构", "4.0", "专业必修", "在修", "2026-2027-1", None)],
    )
    payload = _plan(
        client,
        _import(client, RECORDS_ENDPOINT, records, mime=XLSX_MIME),
        _upload_rule_set(client, tmp_path),
    )
    warning = _warnings(payload)[WARN_TIME_CONFLICT]
    evidence = {item["chunk_id"]: item for item in payload["evidence"]}
    quotes = "".join(evidence[chunk_id]["quote"] for chunk_id in warning["evidence_chunk_ids"])
    assert "QM-GE101" in quotes and "QM-CS201" in quotes


@pytest.mark.parametrize("status", ["通过", "不及格"])
def test_passed_or_failed_schedules_never_conflict(client, db, tmp_path, status: str) -> None:
    """已通过 / 不及格的课程即使带了相同 schedule，也不构成当前选课时间冲突。"""
    records = _write_records_xlsx(
        tmp_path,
        rows=[
            (1, "QM-CS101", "程序设计基础", "4.0", "专业必修", status, "2026-2027-1", "星期二第3-4节"),
            (2, "QM-CS102", "高等数学（一）", "5.0", "公共必修", status, "2026-2027-1", "星期二第3-4节"),
        ],
    )
    payload = _plan(
        client,
        _import(client, RECORDS_ENDPOINT, records, mime=XLSX_MIME),
        _upload_rule_set(client, tmp_path),
    )
    assert WARN_TIME_CONFLICT not in _warnings(payload)


def test_different_semesters_never_conflict(client, db, tmp_path) -> None:
    """两门在修课程的学期不同，即使星期与时间完全相同也不得冲突。"""
    records = _write_records_xlsx(
        tmp_path,
        rows=[
            (1, "QM-CS101", "程序设计基础", "4.0", "专业必修", "在修", "2026-2027-1", "星期二第3-4节"),
            (2, "QM-CS102", "高等数学（一）", "5.0", "公共必修", "在修", "2027-2028-1", "星期二第3-4节"),
        ],
    )
    payload = _plan(
        client,
        _import(client, RECORDS_ENDPOINT, records, mime=XLSX_MIME),
        _upload_rule_set(client, tmp_path),
    )
    assert WARN_TIME_CONFLICT not in _warnings(payload)


def test_missing_semester_never_conflicts(client, db, tmp_path) -> None:
    """学期缺失时无法可靠隔离，不得猜测冲突。"""
    records = _write_records_xlsx(
        tmp_path,
        rows=[
            (1, "QM-CS101", "程序设计基础", "4.0", "专业必修", "在修", None, "星期二第3-4节"),
            (2, "QM-CS102", "高等数学（一）", "5.0", "公共必修", "在修", None, "星期二第3-4节"),
        ],
    )
    payload = _plan(
        client,
        _import(client, RECORDS_ENDPOINT, records, mime=XLSX_MIME),
        _upload_rule_set(client, tmp_path),
    )
    assert WARN_TIME_CONFLICT not in _warnings(payload)


# ---------------------------------------------------------------------------
# 区间边界
# ---------------------------------------------------------------------------


def test_touching_clock_intervals_do_not_conflict(client, db, tmp_path) -> None:
    """09:00-10:00 与 10:00-11:00 首尾相接，按半开区间不算重叠。"""
    records = _write_records_xlsx(
        tmp_path,
        rows=[
            (1, "QM-CS101", "程序设计基础", "4.0", "专业必修", "在修", "2026-2027-1", "星期二 09:00-10:00"),
            (2, "QM-CS102", "高等数学（一）", "5.0", "公共必修", "在修", "2026-2027-1", "星期二 10:00-11:00"),
        ],
    )
    payload = _plan(
        client,
        _import(client, RECORDS_ENDPOINT, records, mime=XLSX_MIME),
        _upload_rule_set(client, tmp_path),
    )
    assert WARN_TIME_CONFLICT not in _warnings(payload)


def test_overlapping_clock_intervals_conflict(client, db, tmp_path) -> None:
    """09:00-10:30 与 10:00-11:00 真实重叠，必须冲突且证据覆盖双方。"""
    records = _write_records_xlsx(
        tmp_path,
        rows=[
            (1, "QM-CS101", "程序设计基础", "4.0", "专业必修", "在修", "2026-2027-1", "星期二 09:00-10:30"),
            (2, "QM-CS102", "高等数学（一）", "5.0", "公共必修", "在修", "2026-2027-1", "星期二 10:00-11:00"),
        ],
    )
    payload = _plan(
        client,
        _import(client, RECORDS_ENDPOINT, records, mime=XLSX_MIME),
        _upload_rule_set(client, tmp_path),
    )
    warning = _warnings(payload)[WARN_TIME_CONFLICT]
    evidence = {item["chunk_id"]: item for item in payload["evidence"]}
    quotes = "".join(evidence[chunk_id]["quote"] for chunk_id in warning["evidence_chunk_ids"])
    assert "QM-CS101" in quotes and "QM-CS102" in quotes


def test_shared_period_is_a_conflict(client, db, tmp_path) -> None:
    """节次是端点包含的离散区间：第1-2节与第2-3节共享第2节，必须冲突。"""
    records = _write_records_xlsx(
        tmp_path,
        rows=[
            (1, "QM-CS101", "程序设计基础", "4.0", "专业必修", "在修", "2026-2027-1", "星期二第1-2节"),
            (2, "QM-CS102", "高等数学（一）", "5.0", "公共必修", "在修", "2026-2027-1", "星期二第2-3节"),
        ],
    )
    payload = _plan(
        client,
        _import(client, RECORDS_ENDPOINT, records, mime=XLSX_MIME),
        _upload_rule_set(client, tmp_path),
    )
    assert WARN_TIME_CONFLICT in _warnings(payload)


# ---------------------------------------------------------------------------
# demo 真实冲突必须保留
# ---------------------------------------------------------------------------


@pytest.mark.parametrize("student", ["匿名学生A · 课程记录", "匿名学生B · 课程记录"])
def test_demo_real_conflict_is_still_detected(client, demo_ready, student: str) -> None:
    records, rules = _demo_ids(client)
    payload = _plan(client, records[student], rules["2026.1"])

    warning = _warnings(payload)[WARN_TIME_CONFLICT]
    assert warning["message"]
    evidence = {item["chunk_id"]: item for item in payload["evidence"]}
    quotes = "".join(evidence[chunk_id]["quote"] for chunk_id in warning["evidence_chunk_ids"])
    assert "QM-GE101" in quotes and "QM-CS201" in quotes, "证据必须覆盖冲突双方课程来源"
