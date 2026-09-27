"""解析器：三种格式的定位保留、稳定顺序与边界场景。"""

from __future__ import annotations

from pathlib import Path

import docx
import openpyxl
import pytest

from tests.conftest import DEMO_DATASET_PATH
from app import constants
from app.demo.manifest import load_manifest
from app.documents.parsing import parse_document

FICTION_MARKER = "仅供系统演示的虚构资料"


def _make_xlsx(path: Path, builder) -> Path:
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    builder(workbook)
    workbook.save(path)
    return path


def _make_docx(path: Path, builder) -> Path:
    document = docx.Document()
    builder(document)
    document.save(path)
    return path


# --- 固化演示语料 -----------------------------------------------------------


def test_all_demo_documents_parse_into_traceable_blocks() -> None:
    manifest = load_manifest(DEMO_DATASET_PATH, "2026.1")
    assert len(manifest.documents) == 15

    for document in manifest.documents:
        path = DEMO_DATASET_PATH / document.path
        blocks = parse_document(path, document.file_type)
        assert blocks, document.path
        assert all(block.text.strip() for block in blocks), document.path
        assert any(FICTION_MARKER in block.text for block in blocks), document.path


def test_parsing_is_stable_across_runs() -> None:
    manifest = load_manifest(DEMO_DATASET_PATH, "2026.1")
    for document in manifest.documents:
        path = DEMO_DATASET_PATH / document.path
        first = parse_document(path, document.file_type)
        second = parse_document(path, document.file_type)
        assert first == second, document.path


def test_pdf_keeps_one_based_page_numbers_and_sections() -> None:
    manifest = load_manifest(DEMO_DATASET_PATH, "2026.1")
    pdf_documents = [item for item in manifest.documents if item.file_type == constants.FILE_TYPE_PDF]
    assert len(pdf_documents) == 5

    for document in pdf_documents:
        blocks = parse_document(DEMO_DATASET_PATH / document.path, "pdf")
        pages = [block.page_number for block in blocks]
        assert all(isinstance(page, int) and page >= 1 for page in pages), document.path
        assert min(pages) == 1
        assert any(block.block_type == "heading" for block in blocks), document.path
        headed = [block for block in blocks if block.section_title]
        assert headed, document.path


def test_docx_keeps_section_titles_and_table_rows() -> None:
    manifest = load_manifest(DEMO_DATASET_PATH, "2026.1")
    docx_documents = [
        item for item in manifest.documents if item.file_type == constants.FILE_TYPE_DOCX
    ]
    assert len(docx_documents) == 5

    for document in docx_documents:
        blocks = parse_document(DEMO_DATASET_PATH / document.path, "docx")
        assert any(block.block_type == "heading" for block in blocks), document.path
        assert any(block.section_title for block in blocks), document.path
        assert all(block.page_number is None for block in blocks), document.path
        # 课程大纲类文档正文包含表格；制度类文档为纯段落文本
        if document.doc_category == "course_syllabus":
            assert any(block.block_type == "table_row" for block in blocks), document.path


def test_xlsx_keeps_sheet_names_and_one_based_rows() -> None:
    manifest = load_manifest(DEMO_DATASET_PATH, "2026.1")
    xlsx_documents = [
        item for item in manifest.documents if item.file_type == constants.FILE_TYPE_XLSX
    ]
    assert len(xlsx_documents) == 5

    for document in xlsx_documents:
        blocks = parse_document(DEMO_DATASET_PATH / document.path, "xlsx")
        assert any(block.sheet_name for block in blocks), document.path
        for block in blocks:
            assert block.row_start is None or block.row_start >= 1, document.path
            assert block.row_end is None or block.row_end >= block.row_start, document.path
            assert block.sheet_name is not None, document.path


# --- XLSX 边界 --------------------------------------------------------------


def test_empty_sheet_is_supported(tmp_path) -> None:
    path = _make_xlsx(
        tmp_path / "mixed.xlsx",
        lambda workbook: (
            workbook.create_sheet("空表"),
            _fill(workbook.create_sheet("数据"), [["事件", "日期"], ["开学", "2026-09-01"]]),
        ),
    )
    blocks = parse_document(path, "xlsx")
    assert {block.sheet_name for block in blocks} == {"数据"}
    assert any("开学" in block.text for block in blocks)


def _fill(sheet, rows):
    for row in rows:
        sheet.append(row)
    return sheet


def test_sheet_caption_row_before_header_is_preserved(tmp_path) -> None:
    path = _make_xlsx(
        tmp_path / "caption.xlsx",
        lambda workbook: _fill(
            workbook.create_sheet("校历"),
            [[FICTION_MARKER], ["事件", "开始日期"], ["开课", "2026-09-01"]],
        ),
    )
    blocks = parse_document(path, "xlsx")
    assert blocks[0].block_type == "caption"
    assert FICTION_MARKER in blocks[0].text
    assert blocks[1].block_type == "table_header"
    assert blocks[1].row_start == 2


def test_merged_cells_are_expanded(tmp_path) -> None:
    def build(workbook):
        sheet = workbook.create_sheet("合并")
        sheet["A1"] = "类别"
        sheet.merge_cells("A1:B1")
        sheet["C1"] = "课程"
        sheet["A2"] = "必修"
        sheet.merge_cells("A2:A3")
        sheet["B2"] = "语文"
        sheet["B3"] = "数学"
        sheet["C2"] = "QM-CS101"
        sheet["C3"] = "QM-CS102"

    path = _make_xlsx(tmp_path / "merged.xlsx", build)
    blocks = parse_document(path, "xlsx")
    header = next(block for block in blocks if block.block_type == "table_header")
    assert "类别" in header.text
    rows = [block.text for block in blocks if block.block_type == "table_row"]
    assert len(rows) == 2
    assert all("必修" in text for text in rows)


def test_formulas_are_not_executed(tmp_path) -> None:
    path = _make_xlsx(
        tmp_path / "formula.xlsx",
        lambda workbook: _fill(
            workbook.create_sheet("公式"),
            [["项目", "数值"], ["合计", "=SUM(1,2)"]],
        ),
    )
    blocks = parse_document(path, "xlsx")
    text = "\n".join(block.text for block in blocks)
    assert "=SUM(1,2)" in text
    assert "数值: 3" not in text


# --- DOCX 边界 --------------------------------------------------------------


def test_docx_paragraph_and_table_order_is_preserved(tmp_path) -> None:
    def build(document):
        document.add_heading("第一章 总则", level=1)
        document.add_paragraph("第一条 说明文字。")
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "项目"
        table.cell(0, 1).text = "内容"
        table.cell(1, 0).text = "学分"
        table.cell(1, 1).text = "4.0"
        document.add_heading("第二章 附则", level=1)
        document.add_paragraph("第二条 结尾段落。")

    path = _make_docx(tmp_path / "order.docx", build)
    blocks = parse_document(path, "docx")
    kinds = [block.block_type for block in blocks]
    assert kinds.index("table_row") < kinds.index("heading") or kinds.count("heading") == 2
    assert [block.text for block in blocks if block.block_type == "heading"] == [
        "第一章 总则",
        "第二章 附则",
    ]
    table_row = next(block for block in blocks if block.block_type == "table_row")
    assert table_row.section_title == "第一章 总则"
    final_paragraph = blocks[-1]
    assert final_paragraph.text == "第二条 结尾段落。"
    assert final_paragraph.section_title == "第二章 附则"


def test_docx_empty_document_is_rejected(tmp_path) -> None:
    from app.core.errors import ApiError

    path = _make_docx(tmp_path / "empty.docx", lambda document: None)
    with pytest.raises(ApiError) as error:
        parse_document(path, "docx")
    assert error.value.code == "DOCUMENT_EMPTY"


def test_unsupported_file_type_is_rejected(tmp_path) -> None:
    from app.core.errors import ApiError

    with pytest.raises(ApiError) as error:
        parse_document(tmp_path / "x.txt", "txt")
    assert error.value.code == "DOCUMENT_INVALID_TYPE"
