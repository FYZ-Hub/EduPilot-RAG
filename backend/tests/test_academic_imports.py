"""阶段 7B-2：学业资料导入接口（records / rules）测试。

全部离线：不访问网络、不调用 LLM / Embedding / Reranker、不下载模型。
导入复用正式的上传安全校验、存储与解析组件；业务字段只来自 DocumentBlock。
"""

from __future__ import annotations

import asyncio
import io
import zipfile
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

import openpyxl
import pytest
from sqlalchemy import func, select
from starlette.datastructures import Headers, UploadFile

from app.academic import imports
from app.models import (
    AcademicRecordSet,
    AcademicRuleSet,
    CourseRecordRow,
    DegreeRuleCourse,
    DegreeRuleRow,
    Document,
    DocumentBlock,
    DocumentChunk,
)
from tests.conftest import demo_file

RECORDS_ENDPOINT = "/api/academic/records/import"
RULES_ENDPOINT = "/api/academic/rules/import"

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


# ---------------------------------------------------------------------------
# 辅助
# ---------------------------------------------------------------------------


def _stored_files(settings) -> list[Path]:
    root = Path(settings.upload_path)
    if not root.exists():
        return []
    return [path for path in root.iterdir() if path.is_file()]


def _post(client, endpoint: str, path: Path, *, filename=None, name=None, mime=None):
    with open(path, "rb") as handle:
        files = {"file": (filename or path.name, handle, mime or "application/octet-stream")}
        data = {} if name is None else {"name": name}
        return client.post(endpoint, files=files, data=data)


def _post_bytes(client, endpoint: str, payload: bytes, *, filename, name=None, mime=None):
    files = {"file": (filename, payload, mime or "application/octet-stream")}
    data = {} if name is None else {"name": name}
    return client.post(endpoint, files=files, data=data)


def _assert_error(response, code: str) -> dict:
    assert response.status_code >= 400, response.text
    payload = response.json()
    assert payload["code"] == code, payload
    assert payload["request_id"], "所有后端生成的错误都必须带非空 request_id"
    assert set(payload) == {"code", "message", "details", "request_id"}
    assert "Traceback" not in response.text
    assert "/app/" not in response.text
    return payload


def _write_records_xlsx(tmp_path: Path, rows, *, header=None) -> Path:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "课程记录"
    sheet.append(list(header or ("序号", "课程代码", "课程名称", "学分", "课程类别", "状态")))
    for row in rows:
        sheet.append(list(row))
    path = tmp_path / "records.xlsx"
    workbook.save(str(path))
    return path


def _write_plan_xlsx(
    tmp_path: Path,
    *,
    fields=None,
    minimums=None,
    catalog=None,
    duplicate_category=False,
) -> Path:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "培养方案"
    sheet.append(["专业名称", "招生年份", "文档版本", "生效日期", "毕业总学分"])
    sheet.append(
        list(fields or ("计算机科学与技术", 2026, "2026.2", "2026-09-01", "160.0 学分"))
    )

    minimum_sheet = workbook.create_sheet("类别最低学分")
    minimum_sheet.append(["课程类别", "最低学分"])
    for category, credits in minimums or (("公共必修", 52), ("专业必修", 58), ("专业选修", 20)):
        minimum_sheet.append([category, credits])
    if duplicate_category:
        minimum_sheet.append(["专业必修", 58])

    catalog_sheet = workbook.create_sheet("课程目录")
    catalog_sheet.append(["课程代码", "课程名称", "学分", "课程类别"])
    for row in catalog or (
        ("QM-CS102", "高等数学（一）", 5, "公共必修"),
        ("QM-CS101", "程序设计基础", 4, "专业必修"),
        ("QM-CS401", "软件工程", 3, "专业选修"),
    ):
        catalog_sheet.append(list(row))

    path = tmp_path / "plan.xlsx"
    workbook.save(str(path))
    return path


def _write_plan_docx(tmp_path: Path, *, version: str = "2026.2", credits: str = "160.0") -> Path:
    import docx

    document = docx.Document()
    document.add_paragraph("【仅供系统演示的虚构资料】本文件为自动化测试确定性生成。")
    document.add_heading("启明大学计算机科学与技术专业培养方案", level=1)
    document.add_paragraph("专业名称：计算机科学与技术")
    document.add_paragraph("招生年份：2026")
    document.add_paragraph(f"文档版本：{version} 生效日期：2026-09-01")
    document.add_paragraph(f"毕业总学分：{credits} 学分。")
    document.add_paragraph("· 公共必修：52.0 学分")
    document.add_paragraph("· 专业必修：58.0 学分")
    document.add_paragraph("· 专业选修：20.0 学分")
    document.add_paragraph("QM-CS102 高等数学（一） 5.0 公共必修 第1学期 无")
    document.add_paragraph("QM-CS101 程序设计基础 4.0 专业必修 第1学期 无")
    document.add_paragraph("QM-CS401 软件工程 3.0 专业选修 第5学期 无")
    path = tmp_path / "plan.docx"
    document.save(str(path))
    return path


def _set_of(context, model, set_id: str):
    with context.session_factory() as session:
        return session.get(model, set_id)


def _document_of(context, set_id: str, model=AcademicRecordSet) -> Document:
    with context.session_factory() as session:
        record = session.get(model, set_id)
        assert record is not None
        document = session.get(Document, record.source_doc_id)
        assert document is not None
        return document


# ---------------------------------------------------------------------------
# 1/2：课程记录成功导入与响应契约
# ---------------------------------------------------------------------------


def test_import_records_from_demo_xlsx(client, context, settings) -> None:
    path = demo_file("13-课程记录-匿名学生A")
    response = _post(client, RECORDS_ENDPOINT, path, mime=XLSX_MIME)

    assert response.status_code == 200, response.text
    payload = response.json()
    assert set(payload) == {"id", "status", "warnings"}
    assert payload["status"] == "ready"
    assert payload["warnings"] == []

    with context.session_factory() as session:
        rows = list(
            session.scalars(
                select(CourseRecordRow)
                .where(CourseRecordRow.record_set_id == payload["id"])
                .order_by(CourseRecordRow.ordinal)
            )
        )
    assert rows, "必须写入课程记录子项"
    # 「汇总」工作表不得被误投影为课程（只保留 QM- 前缀的真实课程代码）
    assert all(row.course_code.startswith("QM-") for row in rows)
    assert len(_stored_files(settings)) == 1


def test_imported_record_rows_match_document_blocks(client, context) -> None:
    from decimal import Decimal

    from app.academic.types import STATUS_IN_PROGRESS, STATUS_PASSED

    path = demo_file("13-课程记录-匿名学生A")
    set_id = _post(client, RECORDS_ENDPOINT, path, mime=XLSX_MIME).json()["id"]

    with context.session_factory() as session:
        rows = list(
            session.scalars(
                select(CourseRecordRow)
                .where(CourseRecordRow.record_set_id == set_id)
                .order_by(CourseRecordRow.ordinal)
            )
        )
        blocks = session.scalar(
            select(func.count(DocumentBlock.id)).where(
                DocumentBlock.doc_id == session.get(AcademicRecordSet, set_id).source_doc_id
            )
        )
    assert blocks and blocks > 0, "必须建立真实 DocumentBlock 来源"
    passed = [row for row in rows if row.status == STATUS_PASSED]
    ongoing = [row for row in rows if row.status == STATUS_IN_PROGRESS]
    assert sum(Decimal(row.credits) for row in passed) == Decimal("20.5")
    assert sum(Decimal(row.credits) for row in ongoing) == Decimal("9.0")
    assert all(row.status in ("passed", "failed", "in_progress") for row in rows)


def test_import_records_rejects_non_xlsx(client) -> None:
    pdf = demo_file("01-培养方案")
    _assert_error(_post(client, RECORDS_ENDPOINT, pdf), "ACADEMIC_FILE_TYPE_UNSUPPORTED")

    docx = demo_file("03-课程大纲-QM-CS201")
    _assert_error(_post(client, RECORDS_ENDPOINT, docx), "ACADEMIC_FILE_TYPE_UNSUPPORTED")


def test_import_rules_rejects_unsupported_format(client) -> None:
    _assert_error(
        _post_bytes(client, RULES_ENDPOINT, b"hello", filename="notes.txt", mime="text/plain"),
        "DOCUMENT_INVALID_TYPE",
    )


# ---------------------------------------------------------------------------
# 3/4/5：培养方案 PDF / DOCX / XLSX 真实解析导入
# ---------------------------------------------------------------------------


def test_import_rules_pdf(client, context) -> None:
    path = demo_file("01-培养方案")
    response = _post(client, RULES_ENDPOINT, path, mime="application/pdf")
    assert response.status_code == 200, response.text
    payload = response.json()
    assert set(payload) == {"id", "status", "warnings"}
    assert payload["status"] == "ready"

    with context.session_factory() as session:
        rule_set = session.get(AcademicRuleSet, payload["id"])
        assert rule_set is not None
        assert rule_set.major == "计算机科学与技术"
        assert rule_set.admission_year == 2025
        assert rule_set.rule_version == "2025.1"
        assert rule_set.effective_from == "2025-09-01"
        assert rule_set.required_credits == "155.0"
        categories = {
            row.category: row.minimum_credits
            for row in session.scalars(
                select(DegreeRuleRow).where(DegreeRuleRow.rule_set_id == payload["id"])
            )
        }
        assert categories["专业必修"] == "58.0"
        assert categories["公共必修"] == "52.0"
        course_codes = set(
            session.scalars(
                select(DegreeRuleCourse.course_code).where(
                    DegreeRuleCourse.rule_set_id == payload["id"]
                )
            )
        )
    assert {"QM-CS101", "QM-CS102", "QM-CS401"} <= course_codes


def test_import_rules_docx(client, context, tmp_path) -> None:
    path = _write_plan_docx(tmp_path)
    response = _post(client, RULES_ENDPOINT, path, mime=None)
    assert response.status_code == 200, response.text
    payload = response.json()

    with context.session_factory() as session:
        rule_set = session.get(AcademicRuleSet, payload["id"])
        assert rule_set.major == "计算机科学与技术"
        assert rule_set.rule_version == "2026.2"
        assert rule_set.required_credits == "160.0"
        courses = list(
            session.scalars(
                select(DegreeRuleCourse)
                .where(DegreeRuleCourse.rule_set_id == payload["id"])
                .order_by(DegreeRuleCourse.course_code)
            )
        )
    assert [course.course_code for course in courses] == ["QM-CS101", "QM-CS102", "QM-CS401"]
    assert courses[0].credits == "4.0"


def test_import_rules_xlsx(client, context, tmp_path) -> None:
    path = _write_plan_xlsx(tmp_path)
    response = _post(client, RULES_ENDPOINT, path, mime=XLSX_MIME)
    assert response.status_code == 200, response.text
    payload = response.json()

    with context.session_factory() as session:
        rule_set = session.get(AcademicRuleSet, payload["id"])
        assert rule_set.major == "计算机科学与技术"
        assert rule_set.admission_year == 2026
        assert rule_set.rule_version == "2026.2"
        assert rule_set.effective_from == "2026-09-01"
        assert rule_set.required_credits == "160.0"
        declared = {
            row.category: (row.minimum_credits, list(row.required_course_codes))
            for row in session.scalars(
                select(DegreeRuleRow).where(DegreeRuleRow.rule_set_id == payload["id"])
            )
        }
    assert declared["专业必修"] == ("58.0", ["QM-CS101"])
    assert declared["公共必修"] == ("52.0", ["QM-CS102"])
    assert declared["专业选修"] == ("20.0", [])


def test_two_rule_versions_are_both_preserved(client) -> None:
    """同一专业的不同版本必须分别保留，后端不得静默只留最新版。"""
    first = _post(client, RULES_ENDPOINT, demo_file("01-培养方案"), mime="application/pdf")
    second = _post(client, RULES_ENDPOINT, demo_file("02-培养方案"), mime="application/pdf")
    assert first.json()["id"] != second.json()["id"]
    assert first.status_code == second.status_code == 200


# ---------------------------------------------------------------------------
# 6—10：安全校验（复用正式上传安全组件）
# ---------------------------------------------------------------------------


@pytest.mark.parametrize(
    "filename",
    [
        "../evil.xlsx",
        "..\\evil.xlsx",
        "/etc/passwd.xlsx",
        "C:\\windows\\evil.xlsx",
        "\\\\server\\share\\evil.xlsx",
        "folder/evil.xlsx",
        "con.xlsx",
        "nul.xlsx",
        "evil\u202eexe.xlsx",
        "evil\u2215path.xlsx",
        "trailing .xlsx ",
        "onlydots..xlsx",
    ],
)
def test_unsafe_filenames_are_rejected_on_import(client, filename: str) -> None:
    path = demo_file("13-课程记录-匿名学生A")
    for endpoint in (RECORDS_ENDPOINT, RULES_ENDPOINT):
        _assert_error(_post(client, endpoint, path, filename=filename), "DOCUMENT_UNSAFE_NAME")


def test_control_characters_in_filename_are_rejected_on_import(client) -> None:
    boundary = "----campusacademic"
    payload = demo_file("13-课程记录-匿名学生A").read_bytes()
    body = (
        f"--{boundary}\r\n"
        'Content-Disposition: form-data; name="file"; filename="bad\x01name.xlsx"\r\n'
        "Content-Type: application/octet-stream\r\n\r\n"
    ).encode("utf-8") + payload + f"\r\n--{boundary}--\r\n".encode("utf-8")
    response = client.post(
        RECORDS_ENDPOINT,
        content=body,
        headers={"Content-Type": f"multipart/form-data; boundary={boundary}"},
    )
    _assert_error(response, "DOCUMENT_UNSAFE_NAME")


def test_missing_file_field_is_rejected(client) -> None:
    response = client.post(RECORDS_ENDPOINT)
    assert response.status_code == 422
    assert response.json()["request_id"]


def test_mime_and_extension_mismatch_are_rejected_on_import(client) -> None:
    path = demo_file("13-课程记录-匿名学生A")
    _assert_error(
        _post(client, RECORDS_ENDPOINT, path, filename="records.xlsx", mime="text/plain"),
        "DOCUMENT_CONTENT_TYPE_MISMATCH",
    )
    _assert_error(
        _post(client, RECORDS_ENDPOINT, path, filename="records.docx", mime=XLSX_MIME),
        "DOCUMENT_CONTENT_TYPE_MISMATCH",
    )


def test_signature_mismatch_is_rejected_on_import(client) -> None:
    """真实的 PDF 字节改成 .xlsx 必须被拒绝（文件头 / 容器不一致）。"""
    response = _post_bytes(
        client,
        RULES_ENDPOINT,
        demo_file("01-培养方案").read_bytes(),
        filename="fake.xlsx",
        mime=XLSX_MIME,
    )
    _assert_error(response, "DOCUMENT_CORRUPT")


def _zip_bytes(entries: dict[str, bytes]) -> bytes:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, payload in entries.items():
            archive.writestr(name, payload)
    return buffer.getvalue()


XLSX_CONTENT_TYPES = (
    '<?xml version="1.0" encoding="UTF-8"?>'
    '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
    '<Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-'
    'officedocument.spreadsheetml.sheet.main+xml"/></Types>'
)


def test_zip_traversal_macro_and_bomb_are_rejected_on_import(client) -> None:
    traversal = _zip_bytes(
        {"[Content_Types].xml": XLSX_CONTENT_TYPES.encode(), "../evil.xml": b"x"}
    )
    _assert_error(
        _post_bytes(client, RULES_ENDPOINT, traversal, filename="t.xlsx", mime=XLSX_MIME),
        "DOCUMENT_UNSAFE_CONTAINER",
    )

    macro = _zip_bytes(
        {"[Content_Types].xml": XLSX_CONTENT_TYPES.encode(), "xl/vbaProject.bin": b"macro"}
    )
    _assert_error(
        _post_bytes(client, RULES_ENDPOINT, macro, filename="m.xlsx", mime=XLSX_MIME),
        "DOCUMENT_UNSAFE_CONTAINER",
    )

    bomb = _zip_bytes({"payload.bin": b"\0" * (2 * 1024 * 1024)})
    _assert_error(
        _post_bytes(client, RULES_ENDPOINT, bomb, filename="b.xlsx", mime=XLSX_MIME),
        "DOCUMENT_UNSAFE_CONTAINER",
    )


def test_oversized_file_is_rejected_on_import(tmp_path) -> None:
    from fastapi.testclient import TestClient

    from app.main import create_app
    from tests.conftest import build_settings

    settings = build_settings(tmp_path, max_upload_mb=1)
    application = create_app(settings)
    with TestClient(application) as test_client:
        payload = demo_file("13-课程记录-匿名学生A").read_bytes() + b"0" * (2 * 1024 * 1024)
        response = _post_bytes(
            test_client, RECORDS_ENDPOINT, payload, filename="big.xlsx", mime=XLSX_MIME
        )
        _assert_error(response, "DOCUMENT_TOO_LARGE")
        assert not _stored_files(settings)
    application.state.context.engine.dispose()


# ---------------------------------------------------------------------------
# 11/12：字段、取值与规则冲突
# ---------------------------------------------------------------------------


def test_missing_required_header_fails(client, tmp_path) -> None:
    path = _write_records_xlsx(
        tmp_path,
        [[1, "QM-CS101", "程序设计基础", "专业必修", "通过"]],
        header=("序号", "课程代码", "课程名称", "课程类别", "状态"),
    )
    _assert_error(
        _post(client, RECORDS_ENDPOINT, path, mime=XLSX_MIME), "ACADEMIC_FIELD_MISSING"
    )


def test_invalid_status_fails(client, tmp_path) -> None:
    path = _write_records_xlsx(
        tmp_path, [[1, "QM-CS101", "程序设计基础", "4.0", "专业必修", "优异"]]
    )
    _assert_error(
        _post(client, RECORDS_ENDPOINT, path, mime=XLSX_MIME), "ACADEMIC_VALUE_INVALID"
    )


@pytest.mark.parametrize("credits", ["-3", "abc", "inf", "nan", "100000"])
def test_invalid_credits_fail(client, tmp_path, credits: str) -> None:
    path = _write_records_xlsx(
        tmp_path, [[1, "QM-CS101", "程序设计基础", credits, "专业必修", "通过"]]
    )
    _assert_error(
        _post(client, RECORDS_ENDPOINT, path, mime=XLSX_MIME), "ACADEMIC_VALUE_INVALID"
    )


def test_duplicate_identical_category_yields_stable_warning(client, tmp_path) -> None:
    path = _write_plan_xlsx(tmp_path, duplicate_category=True)
    response = _post(client, RULES_ENDPOINT, path, mime=XLSX_MIME)
    assert response.status_code == 200, response.text
    assert response.json()["warnings"] == ["DUPLICATE_CATEGORY_DECLARATION"]


def test_conflicting_category_minimum_fails(client, tmp_path) -> None:
    path = _write_plan_xlsx(
        tmp_path,
        minimums=(("公共必修", 52), ("专业必修", 58), ("专业必修", 60)),
    )
    _assert_error(_post(client, RULES_ENDPOINT, path, mime=XLSX_MIME), "ACADEMIC_RULE_CONFLICT")


def test_conflicting_course_code_fails(client, tmp_path) -> None:
    path = _write_plan_xlsx(
        tmp_path,
        catalog=(
            ("QM-CS101", "程序设计基础", 4, "专业必修"),
            ("QM-CS101", "程序设计基础（二）", 4, "专业必修"),
        ),
    )
    _assert_error(_post(client, RULES_ENDPOINT, path, mime=XLSX_MIME), "ACADEMIC_RULE_CONFLICT")


def test_duplicate_identical_catalog_row_yields_stable_warning(client, tmp_path) -> None:
    path = _write_plan_xlsx(
        tmp_path,
        catalog=(
            ("QM-CS101", "程序设计基础", 4, "专业必修"),
            ("QM-CS101", "程序设计基础", 4, "专业必修"),
        ),
    )
    response = _post(client, RULES_ENDPOINT, path, mime=XLSX_MIME)
    assert response.status_code == 200, response.text
    assert response.json()["warnings"] == ["DUPLICATE_CATALOG_ENTRY"]


def test_missing_plan_field_fails(client, tmp_path) -> None:
    path = _write_plan_xlsx(tmp_path)
    # 去掉「毕业总学分」表头列，必须明确失败而不是猜默认值
    workbook = openpyxl.load_workbook(str(path))
    sheet = workbook["培养方案"]
    sheet.cell(row=1, column=5).value = "备注"
    workbook.save(str(path))
    _assert_error(_post(client, RULES_ENDPOINT, path, mime=XLSX_MIME), "ACADEMIC_FIELD_MISSING")


# ---------------------------------------------------------------------------
# 13—18：幂等、并发、原子性与来源真实性
# ---------------------------------------------------------------------------


def _counts(context) -> dict:
    with context.session_factory() as session:
        return {
            "documents": session.scalar(select(func.count(Document.id))),
            "blocks": session.scalar(select(func.count(DocumentBlock.id))),
            "record_sets": session.scalar(select(func.count(AcademicRecordSet.id))),
            "course_records": session.scalar(select(func.count(CourseRecordRow.id))),
        }


def test_repeated_import_returns_same_set(client, context, settings) -> None:
    path = demo_file("13-课程记录-匿名学生A")
    first = _post(client, RECORDS_ENDPOINT, path, mime=XLSX_MIME)
    before = _counts(context)

    second = _post(client, RECORDS_ENDPOINT, path, mime=XLSX_MIME)
    assert first.json() == second.json()
    assert _counts(context) == before
    assert len(_stored_files(settings)) == 1


def test_repeated_import_with_different_name_is_still_idempotent(
    client, context, settings
) -> None:
    path = demo_file("13-课程记录-匿名学生A")
    first = _post(client, RECORDS_ENDPOINT, path, name="我的成绩单", mime=XLSX_MIME)
    before = _counts(context)

    second = _post(client, RECORDS_ENDPOINT, path, name="换一个名字", mime=XLSX_MIME)
    assert second.json()["id"] == first.json()["id"]
    assert _counts(context) == before
    assert len(_stored_files(settings)) == 1
    # 改显示名不能绕过内容幂等：集合名保持首次导入的结果
    assert _set_of(context, AcademicRecordSet, first.json()["id"]).display_name == "我的成绩单"


def test_imported_name_is_validated(client) -> None:
    path = demo_file("13-课程记录-匿名学生A")
    _assert_error(
        _post(client, RECORDS_ENDPOINT, path, name="../../etc/passwd", mime=XLSX_MIME),
        "DOCUMENT_UNSAFE_NAME",
    )
    _assert_error(
        _post(client, RECORDS_ENDPOINT, path, name="bad\x01name", mime=XLSX_MIME),
        "DOCUMENT_UNSAFE_NAME",
    )


def test_default_display_name_comes_from_safe_file_name(client, context) -> None:
    path = demo_file("13-课程记录-匿名学生A")
    set_id = _post(
        client, RECORDS_ENDPOINT, path, filename="我的课程记录.xlsx", mime=XLSX_MIME
    ).json()["id"]
    record = _set_of(context, AcademicRecordSet, set_id)
    assert "/" not in record.display_name and "\\" not in record.display_name
    assert "C:" not in record.display_name and "/app" not in record.display_name


def test_identity_columns_never_reach_records_or_options(client, context, tmp_path) -> None:
    """记录表里的额外身份列不得进入 CourseRecord 或 options。"""
    path = _write_records_xlsx(
        tmp_path,
        [[1, "QM-CS101", "程序设计基础", "4.0", "专业必修", "通过", "张三", "2026000001"]],
        header=(
            "序号",
            "课程代码",
            "课程名称",
            "学分",
            "课程类别",
            "状态",
            "姓名",
            "学号",
        ),
    )
    set_id = _post(client, RECORDS_ENDPOINT, path, mime=XLSX_MIME).json()["id"]

    with context.session_factory() as session:
        record = session.get(AcademicRecordSet, set_id)
        rows = list(
            session.scalars(select(CourseRecordRow).where(CourseRecordRow.record_set_id == set_id))
        )
    rendered = repr(rows) + record.display_name
    assert "张三" not in rendered
    assert "2026000001" not in rendered

    options_text = client.get("/api/academic/options").text
    assert "张三" not in options_text and "2026000001" not in options_text


def _import_in_thread(settings, session_factory, payload: bytes, filename: str) -> str:
    async def runner() -> str:
        upload = UploadFile(
            file=io.BytesIO(payload),
            filename=filename,
            headers=Headers({"content-type": "application/octet-stream"}),
        )
        with session_factory() as session:
            result = await imports.import_record_set(
                upload, name=None, settings=settings, session=session
            )
            return result.set_id

    return asyncio.run(runner())


def test_concurrent_import_creates_single_set(client, context, settings) -> None:
    payload = demo_file("13-课程记录-匿名学生A").read_bytes()
    with ThreadPoolExecutor(max_workers=2) as pool:
        futures = [
            pool.submit(_import_in_thread, settings, context.session_factory, payload, "records.xlsx")
            for _ in range(2)
        ]
        ids = [future.result() for future in futures]

    assert ids[0] == ids[1], "并发相同导入必须返回同一 ID"
    assert _counts(context)["record_sets"] == 1
    assert _counts(context)["documents"] == 1
    assert len(_stored_files(settings)) == 1


def test_failed_import_leaves_no_partial_rows_or_files(client, context, settings, tmp_path) -> None:
    path = _write_records_xlsx(
        tmp_path, [[1, "QM-CS101", "程序设计基础", "-1", "专业必修", "通过"]]
    )
    before = _counts(context)
    _assert_error(
        _post(client, RECORDS_ENDPOINT, path, mime=XLSX_MIME), "ACADEMIC_VALUE_INVALID"
    )
    assert _counts(context) == before
    assert _stored_files(settings) == []
    temp_dir = Path(settings.upload_tmp_path)
    assert not temp_dir.exists() or not list(temp_dir.glob("*"))


def test_failed_rule_import_leaves_no_partial_rows(client, context, tmp_path) -> None:
    path = _write_plan_xlsx(
        tmp_path, catalog=(("QM-CS101", "程序设计基础", 4, "专业必修"),)
    )
    workbook = openpyxl.load_workbook(str(path))
    workbook["课程目录"].cell(row=2, column=4).value = "不存在类别"
    workbook.save(str(path))

    with context.session_factory() as session:
        before = session.scalar(select(func.count(AcademicRuleSet.id)))
    _assert_error(_post(client, RULES_ENDPOINT, path, mime=XLSX_MIME), "ACADEMIC_RULE_CONFLICT")
    with context.session_factory() as session:
        assert session.scalar(select(func.count(AcademicRuleSet.id))) == before
        assert session.scalar(select(func.count(DegreeRuleRow.id))) == 0
        assert session.scalar(select(func.count(DegreeRuleCourse.id))) == 0


def test_source_ids_are_real_and_same_document(client, context) -> None:
    path = demo_file("13-课程记录-匿名学生A")
    set_id = _post(client, RECORDS_ENDPOINT, path, mime=XLSX_MIME).json()["id"]

    with context.session_factory() as session:
        record = session.get(AcademicRecordSet, set_id)
        document = session.get(Document, record.source_doc_id)
        assert document is not None and document.deleted_at is None
        rows = list(
            session.scalars(select(CourseRecordRow).where(CourseRecordRow.record_set_id == set_id))
        )
        assert rows
        for row in rows:
            assert row.source_doc_id == document.id
            # 学业导入不建立切片，因此必须留空而不是伪造 chunk_id
            assert row.source_chunk_id is None
        chunk_count = session.scalar(
            select(func.count(DocumentChunk.id)).where(DocumentChunk.doc_id == document.id)
        )
    assert chunk_count == 0


# ---------------------------------------------------------------------------
# 19：RAG 检索隔离
# ---------------------------------------------------------------------------


def test_imported_document_stays_out_of_rag(client, context, worker) -> None:
    path = demo_file("13-课程记录-匿名学生A")
    set_id = _post(client, RECORDS_ENDPOINT, path, mime=XLSX_MIME).json()["id"]
    document = _document_of(context, set_id)

    assert document.retrievable is False
    assert document.status == "queued"
    assert document.current_stage is None
    assert document.activation_state is None

    # worker 不得认领学业导入文档（否则会把它变成可检索语料）
    assert worker.run_once() is False

    listing = client.get("/api/documents").json()
    entry = next(item for item in listing["items"] if item["id"] == document.id)
    assert entry["retrievable"] is False
    assert listing["counts"]["retrievable"] == 0

    options = client.get("/api/retrieval/options").json()
    assert options["doc_categories"] == []

    with context.session_factory() as session:
        refreshed = session.get(Document, document.id)
    assert refreshed.retrievable is False
    assert refreshed.status == "queued"


def test_source_commit_failure_returns_stable_error(client, monkeypatch) -> None:
    """原始文件无法落盘时不得谎报成功，必须返回稳定错误且不留下文件。"""
    from app.academic import imports as imports_module

    def _boom(pending, settings):
        raise OSError("disk full")

    monkeypatch.setattr(imports_module, "commit_upload", _boom)
    response = _post(client, RECORDS_ENDPOINT, demo_file("13-课程记录-匿名学生A"), mime=XLSX_MIME)
    _assert_error(response, "ACADEMIC_SOURCE_UNAVAILABLE")


def test_health_still_reports_planning_unavailable(client) -> None:
    payload = client.get("/api/health").json()
    assert payload["capabilities"]["planning"] == "unavailable"


def test_plan_endpoint_is_not_implemented(client) -> None:
    response = client.post("/api/academic/plan", json={})
    assert response.status_code == 404
    assert response.json()["code"] == "NOT_FOUND"
