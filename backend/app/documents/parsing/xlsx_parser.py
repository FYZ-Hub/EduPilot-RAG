"""XLSX 解析（openpyxl），保留工作表、表头与 1-based 行范围；公式不执行。"""

from __future__ import annotations

from pathlib import Path
from typing import Any

import openpyxl

from app.core.errors import DOCUMENT_CORRUPT, ApiError
from app.documents.blocks import ParsedBlock, cell_text, key_value_row

# 防御异常大的工作表：超过上限的尾部行/列不再解析
MAX_SHEET_ROWS = 5000
MAX_SHEET_COLUMNS = 100


def parse_xlsx(path: Path) -> list[ParsedBlock]:
    """解析全部工作表；空工作表不产生块，也不视为错误。"""
    try:
        workbook = openpyxl.load_workbook(str(path), data_only=False, read_only=False)
    except Exception as error:  # noqa: BLE001 - 第三方解析异常统一收敛
        raise ApiError(DOCUMENT_CORRUPT, retryable=False) from error

    blocks: list[ParsedBlock] = []
    try:
        for sheet_index, worksheet in enumerate(workbook.worksheets):
            blocks.extend(_sheet_blocks(worksheet, sheet_index))
    finally:
        workbook.close()
    return blocks


def _merged_anchors(worksheet) -> dict[tuple[int, int], Any]:
    """把合并区域的左上角取值映射到区域内其它坐标，避免出现空洞。"""
    anchors: dict[tuple[int, int], Any] = {}
    for merged in worksheet.merged_cells.ranges:
        anchor = worksheet.cell(row=merged.min_row, column=merged.min_col).value
        if anchor is None:
            continue
        for row in range(merged.min_row, merged.max_row + 1):
            for column in range(merged.min_col, merged.max_col + 1):
                if (row, column) == (merged.min_row, merged.min_col):
                    continue
                anchors[(row, column)] = anchor
    return anchors


def _sheet_blocks(worksheet, sheet_index: int) -> list[ParsedBlock]:
    max_row = min(worksheet.max_row or 0, MAX_SHEET_ROWS)
    max_column = min(worksheet.max_column or 0, MAX_SHEET_COLUMNS)
    if max_row == 0 or max_column == 0:
        return []

    merged = _merged_anchors(worksheet)
    rows: list[tuple[int, list[str]]] = []
    for row_index in range(1, max_row + 1):
        values: list[str] = []
        for column in range(1, max_column + 1):
            value = worksheet.cell(row=row_index, column=column).value
            if value is None:
                value = merged.get((row_index, column))
            values.append(cell_text(value))
        rows.append((row_index, values))

    header_position = _header_position(rows)
    if header_position is None:
        return []

    # 表头之前的非空行按说明文字处理（例如工作表顶部的资料标识行）
    blocks: list[ParsedBlock] = []
    for row_number, values in rows[:header_position]:
        text = " | ".join(value for value in values if value)
        if not text:
            continue
        blocks.append(
            ParsedBlock(
                text=text,
                block_type="caption",
                sheet_name=worksheet.title,
                row_start=row_number,
                row_end=row_number,
                metadata={"sheet_index": sheet_index},
            )
        )

    header_row_number, raw_header = rows[header_position]
    header = [name or f"第{index + 1}列" for index, name in enumerate(raw_header)]

    blocks.append(
        ParsedBlock(
            text=" | ".join(header),
            block_type="table_header",
            sheet_name=worksheet.title,
            row_start=header_row_number,
            row_end=header_row_number,
            metadata={"sheet_index": sheet_index, "header": header},
        )
    )
    for row_number, values in rows[header_position + 1:]:
        if not any(values):
            continue
        text = key_value_row(header, values) or " | ".join(values)
        blocks.append(
            ParsedBlock(
                text=text,
                block_type="table_row",
                sheet_name=worksheet.title,
                row_start=row_number,
                row_end=row_number,
                metadata={"sheet_index": sheet_index, "header": header, "cells": values},
            )
        )
    return blocks


def _header_position(rows: list[tuple[int, list[str]]]) -> int | None:
    """表头取第一行“至少两个非空单元格”的行；否则退回第一行非空行。"""
    for index, (_, values) in enumerate(rows):
        if sum(1 for value in values if value) >= 2:
            return index
    return next((index for index, (_, values) in enumerate(rows) if any(values)), None)
